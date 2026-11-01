"""Extract local task evidence with locators. Never execute documents or call models/network APIs."""
import csv
import contextlib
import datetime
import io
import json
import math
import pathlib
import re
import struct
import sys
import zipfile

from extract_artifact import extract as extract_legacy

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_SEGMENTS = 2000
MAX_TEXT = 90000


class LimitError(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, message


class Result:
    def __init__(self, format_name):
        self.data = {"version": 1, "format": format_name, "segments": [], "issues": []}
        self.characters = 0

    def issue(self, code, severity, message, segment_id=None):
        issue = {"code": code, "severity": severity, "message": message}
        if segment_id is not None:
            issue["segmentId"] = segment_id
        self.data["issues"].append(issue)

    def add(self, segment_id, text, locator, fields=None):
        if len(self.data["segments"]) >= MAX_SEGMENTS:
            raise LimitError("segment_limit", "Source exceeds 2000 segments. Split the selected work item explicitly.")
        self.characters += len(text)
        if self.characters > MAX_TEXT:
            raise LimitError("content_limit", "Extracted content exceeds 90000 characters. No partial extraction was retained.")
        segment = {"id": segment_id, "text": text, "locator": locator}
        if fields is not None:
            segment["fields"] = fields
        self.data["segments"].append(segment)


def literal(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("non_finite_number")
    return value


def display(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(literal(value))


def adf_text(value):
    """Flatten only supplied ADF text nodes; do not infer links, diagrams or criteria."""
    parts = []

    def walk(item):
        if isinstance(item, dict):
            if isinstance(item.get("text"), str):
                parts.append(item["text"])
            for child in item.get("content", []):
                walk(child)
        elif isinstance(item, list):
            for child in item:
                walk(child)

    walk(value)
    return " ".join(part.strip() for part in parts if part.strip())


def readable(value, prefix=""):
    if isinstance(value, dict):
        if value.get("type") == "doc" and isinstance(value.get("content"), list):
            return (prefix + ": " if prefix else "") + adf_text(value)
        return "\n".join(readable(child, f"{prefix}.{key}" if prefix else str(key)) for key, child in value.items())
    if isinstance(value, list):
        return "\n".join(readable(child, f"{prefix}[{index}]") for index, child in enumerate(value))
    return (prefix + ": " if prefix else "") + display(value)


def column_name(index):
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def table(result, rows, sheet=None, formulas=None, row_numbers=None):
    rows = iter(rows)
    headers = None
    for physical_row, values in enumerate(rows, 1):
        if physical_row > MAX_SEGMENTS:
            raise LimitError("row_limit", "A table exceeds 2000 rows, including its header. Split the selected sheet.")
        values = [literal(value) for value in values]
        if not any(value is not None and display(value).strip() for value in values):
            continue
        row = row_numbers[physical_row - 1] if row_numbers else physical_row
        segment_id = ("sheet-" + str(sheet[0]) + "-" if sheet else "") + "row-" + str(row)
        location = {"row": row, "cells": f"A{row}:{column_name(len(values))}{row}"}
        if sheet:
            location["sheet"] = sheet[1]
        header_row = headers is None
        if header_row:
            headers = [display(value).strip() for value in values]
            location["section"] = "header"
            if any(not name for name in headers) or len(set(headers)) != len(headers):
                result.issue("ambiguous_headers", "blocking", "Table headers are blank or duplicated. Supply an explicit, unambiguous column mapping.", segment_id)
        fields = {}
        for index, value in enumerate(values):
            key = headers[index] if index < len(headers) and headers[index] else "column_" + column_name(index + 1)
            # Duplicate headers cannot be represented safely as one mapped field.
            if key in fields:
                key = "column_" + column_name(index + 1)
                while key in fields:
                    key += "_duplicate"
            if formulas and (row - 1, index) in formulas:
                result.issue("formula_requires_literal", "blocking", f"Formula cell {column_name(index + 1)}{row} must be replaced with a reviewed literal value before acceptance review.", segment_id)
            fields[key] = value
        text = "\n".join(f"{key}: {display(value)}" for key, value in fields.items())
        result.add(segment_id, text, location, None if header_row else fields)


def extract_csv(file, result):
    # csv.reader handles quoted delimiters and embedded newlines; line_num locates record endings.
    with file.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, strict=True)
        rows, lines, previous_line = [], [], 0
        for row in reader:
            rows.append(row)
            lines.append(previous_line + 1)
            previous_line = reader.line_num
            if len(rows) > MAX_SEGMENTS:
                raise LimitError("row_limit", "CSV exceeds 2000 records, including its header.")
        table(result, rows)
        for segment in result.data["segments"]:
            segment["locator"]["line"] = lines[segment["locator"]["row"] - 1]


def extract_json(file, result):
    def invalid_constant(_):
        raise ValueError("non_finite_json_number")

    def unique_fields(pairs):
        value = {}
        for key, item in pairs:
            if key in value:
                raise ValueError("duplicate_json_field")
            value[key] = item
        return value

    value = json.loads(file.read_text(encoding="utf-8-sig"), parse_constant=invalid_constant, object_pairs_hook=unique_fields)
    section, records = "$", None
    if isinstance(value, dict):
        for key in ("requirements", "issues", "tickets"):
            if isinstance(value.get(key), list):
                section, records = key, value[key]
                break
    if records is None:
        records = value if isinstance(value, list) else [value]
    if not records:
        result.issue("empty_records", "blocking", "The selected JSON contains no records.")
    for index, record in enumerate(records):
        result.add("record-" + str(index + 1), readable(record), {"section": f"{section}[{index}]"}, record if isinstance(record, dict) else None)


def extract_lines(text, result, markdown=False, section=None):
    for line, text_line in enumerate(text.splitlines(), 1):
        if markdown:
            heading = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", text_line)
            if heading:
                section = heading.group(1)
        if text_line.strip():
            locator = {"line": line}
            if section:
                locator["section"] = section
            result.add("line-" + str(line), text_line, locator)


def pdf_has_images(page):
    def inspect(resources, depth=0):
        if resources is None or depth > 4:
            return False
        resources = resources.get_object()
        objects = resources.get("/XObject", {})
        objects = objects.get_object() if hasattr(objects, "get_object") else objects
        for item in objects.values():
            obj = item.get_object()
            if obj.get("/Subtype") == "/Image" or inspect(obj.get("/Resources"), depth + 1):
                return True
        return False

    return inspect(page.get("/Resources"))


def extract_pdf(file, result):
    from pypdf import PdfReader
    reader = PdfReader(file)
    if reader.is_encrypted:
        result.issue("encrypted_pdf", "blocking", "Encrypted PDFs require an authorised readable export.")
        return
    if len(reader.pages) > 100:
        raise LimitError("page_limit", "PDF exceeds 100 pages. Select a bounded document or page range explicitly.")
    result.issue("pdf_layout_review", "warning", "PDF text order and visual meaning need review; diagrams and OCR are not interpreted.")
    for index, page in enumerate(reader.pages, 1):
        segment_id = "page-" + str(index)
        text = page.extract_text() or ""
        result.add(segment_id, text, {"page": index})
        if not text.strip():
            result.issue("pdf_page_without_text", "blocking", "This PDF page has no selectable text. Supply reviewed OCR/transcription or mark the page explicitly irrelevant.", segment_id)
        if pdf_has_images(page):
            result.issue("visual_content_review", "warning", "This page contains image content. Review it separately; no image or diagram semantics were inferred.", segment_id)


def extract_xlsx(file, result, selected_sheet=None):
    from openpyxl import load_workbook
    book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
    try:
        if selected_sheet is not None and selected_sheet not in book.sheetnames:
            result.issue("sheet_not_found", "blocking", "Selected worksheet is missing. Choose an exact worksheet name from the original workbook.")
            return
        for sheet_index, sheet in enumerate(book, 1):
            if selected_sheet is not None and sheet.title != selected_sheet:
                continue
            if sheet.max_row and sheet.max_row > MAX_SEGMENTS:
                raise LimitError("row_limit", "A worksheet exceeds 2000 rows, including its header.")
            formulas = set()

            def rows():
                for row_index, row in enumerate(sheet.iter_rows()):
                    values = []
                    for column, cell in enumerate(row):
                        if cell.data_type == "f":
                            formulas.add((row_index, column))
                        values.append(cell.value)
                    yield values

            table(result, rows(), (sheet_index, sheet.title), formulas)
    finally:
        book.close()
    result.issue("spreadsheet_layout_review", "warning", "Cell values are preserved; formatting, charts, merged layouts and inferred traceability require review.")


def xls_formula_cells(book, sheet_index):
    """Read BIFF FORMULA coordinates; xlrd otherwise exposes only cached cell values."""
    from xlrd.biffh import XL_FORMULA_OPCODES
    if book.biff_version < 50:
        raise LimitError("xls_formula_inspection", "This older XLS version cannot be safely checked for formulas. Export XLSX or reviewed literal CSV.")
    position, formulas = book._sh_abs_posn[sheet_index], set()
    memory = book.mem
    while position + 4 <= len(memory):
        code, length = struct.unpack("<HH", memory[position:position + 4])
        position += 4
        if position + length > len(memory):
            raise ValueError("incomplete_biff_record")
        if code in XL_FORMULA_OPCODES:
            if length < 4:
                raise ValueError("incomplete_formula_record")
            formulas.add(struct.unpack("<HH", memory[position:position + 4]))
        position += length
        if code == 0x000A:  # worksheet EOF
            return formulas
    raise ValueError("missing_biff_eof")


def extract_xls(file, result, selected_sheet=None):
    import xlrd
    book = xlrd.open_workbook(str(file), on_demand=True, logfile=io.StringIO())
    try:
        if selected_sheet is not None and selected_sheet not in book.sheet_names():
            result.issue("sheet_not_found", "blocking", "Selected worksheet is missing. Choose an exact worksheet name from the original workbook.")
            return
        for sheet_index in range(book.nsheets):
            if selected_sheet is not None and book.sheet_names()[sheet_index] != selected_sheet:
                continue
            formulas = xls_formula_cells(book, sheet_index)
            sheet = book.sheet_by_index(sheet_index)
            if sheet.nrows > MAX_SEGMENTS:
                raise LimitError("row_limit", "A worksheet exceeds 2000 rows, including its header.")
            def rows():
                for row in range(sheet.nrows):
                    values = []
                    for column in range(sheet.ncols):
                        cell = sheet.cell(row, column)
                        value = bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else cell.value
                        if (row, column) in formulas:
                            value = "[Formula cell; cached value: " + display(value) + "]"
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
                        values.append(value)
                    yield values

            table(result, rows(), (sheet_index + 1, sheet.name), formulas)
    finally:
        book.release_resources()
    result.issue("spreadsheet_layout_review", "warning", "Legacy workbook formatting and cached values need review; detected formula cells are blocked rather than evaluated.")


def extract(file, selected_sheet=None):
    file = pathlib.Path(file)
    ext = file.suffix.lower().lstrip(".")
    result = Result("image" if ext in ("png", "jpg", "jpeg", "gif", "webp") else ext or "unknown")
    if selected_sheet is not None:
        result.data["selection"] = {"sheet": selected_sheet}
    try:
        if file.stat().st_size > MAX_BYTES:
            raise LimitError("file_limit", "Source exceeds 10 MiB. Split the selected document explicitly.")
        if selected_sheet is not None and ext not in ("xls", "xlsx"):
            result.issue("unsupported_sheet_selection", "blocking", "Worksheet selection is supported only for XLS and XLSX sources.")
            return result.data
        if ext in ("docx", "xlsx"):
            with zipfile.ZipFile(file) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise LimitError("expanded_limit", "Expanded Office document exceeds 50 MiB.")
        if result.data["format"] == "image":
            result.issue("visual_input_requires_review", "blocking", "Image evidence needs reviewed transcription and a source reference. No OCR or diagram interpretation was performed.")
        elif ext == "pdf":
            extract_pdf(file, result)
        elif ext == "json":
            extract_json(file, result)
        elif ext == "csv":
            extract_csv(file, result)
        elif ext == "xlsx":
            extract_xlsx(file, result, selected_sheet)
        elif ext == "xls":
            extract_xls(file, result, selected_sheet)
        elif ext in ("md", "txt"):
            extract_lines(file.read_text(encoding="utf-8-sig"), result, ext == "md")
        elif ext in ("doc", "docx", "html", "htm"):
            extract_lines(extract_legacy(file), result, section="extracted document text")
            result.issue("document_layout_review", "warning", "Document text was linearized. Review original tables, images and diagrams; derived line numbers are not original page numbers.")
        else:
            result.issue("unsupported_format", "blocking", "Unsupported input. Supply PDF, Office, HTML, JSON, CSV, Markdown, text or an explicitly reviewed visual input.")
        if not result.data["segments"] and not any(issue["severity"] == "blocking" for issue in result.data["issues"]):
            result.issue("empty_source", "blocking", "No readable source content was extracted.")
    except LimitError as error:
        result.data["segments"] = []
        result.issue(error.code, "blocking", error.message)
    except ModuleNotFoundError:
        result.data["segments"] = []
        result.issue("missing_dependency", "blocking", "Optional document parser is unavailable. Use the configured Python environment with the project's artifact requirements.")
    except Exception:
        result.data["segments"] = []
        result.issue("parse_failed", "blocking", "Source could not be read safely. Check format, permissions and converter availability; parser diagnostics were suppressed.")
    return result.data


if __name__ == "__main__":
    arguments = sys.argv[1:]
    valid = len(arguments) == 1 or len(arguments) == 3 and arguments[1] == "--sheet" and bool(arguments[2].strip())
    if not valid:
        output = Result("unknown")
        output.issue("input_required", "blocking", "Provide one local source file and optional --sheet NAME for a spreadsheet.")
        output = output.data
    else:
        # Third-party parser diagnostics can contain source snippets or filenames.
        # The CLI emits only the reviewed JSON shape, including sanitized issues.
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            output = extract(arguments[0], arguments[2] if len(arguments) == 3 else None)
    print(json.dumps(output, ensure_ascii=False, allow_nan=False))
