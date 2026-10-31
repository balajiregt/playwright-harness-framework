"""Structured local requirements import. No macros, document execution or model calls."""
import json, pathlib, sys, zipfile, xml.etree.ElementTree as ET, subprocess, shutil, tempfile
MAX_BYTES = 10 * 1024 * 1024

def records(rows):
    rows = [r for r in rows if any(str(v or '').strip() for v in r)]
    if not rows: raise ValueError('empty_sheet')
    header = [str(v or '').strip().lower() for v in rows[0]]
    if len(set(header)) != len(header) or not {'id','criterion','journey'}.issubset(header):
        raise ValueError('Required columns: id, criterion, journey, critical (optional)')
    result=[]
    for row in rows[1:]:
        item=dict(zip(header,row))
        result.append(normalize(item))
    return result

def normalize(item):
    if any(k not in {'id','criterion','journey','critical'} for k in item): raise ValueError('Unknown requirement field')
    raw=item.get('critical',False)
    if raw in (None,''): raw=False
    if isinstance(raw,str):
        if raw.strip().lower() not in ('true','false','yes','no','1','0'): raise ValueError('critical must be true or false')
        raw=raw.strip().lower() in ('true','yes','1')
    if raw not in (True,False,0,1): raise ValueError('critical must be true or false')
    result={k:str(item.get(k,'') or '').strip() for k in ('id','criterion','journey')}
    if any(not result[k] for k in result): raise ValueError('Requirement has missing fields')
    result['critical']=bool(raw)
    return result

def blocks(text):
    result=[]; current={}
    for line in text.splitlines():
        line=line.strip()
        if not line: continue
        if ':' not in line: raise ValueError('Use labelled Word blocks: ID:, Criterion:, Journey:, Critical:')
        key,value=line.split(':',1); key=key.strip().lower()
        if key=='id' and current: result.append(normalize(current));current={}
        if key not in ('id','criterion','journey','critical') or key in current: raise ValueError('Unknown or repeated Word field')
        current[key]=value.strip()
    if current: result.append(normalize(current))
    if not result: raise ValueError('No requirements found')
    return result

def load(file):
    file=pathlib.Path(file).resolve()
    if file.stat().st_size>MAX_BYTES: raise ValueError('File exceeds 10 MiB limit')
    ext=file.suffix.lower()
    if ext=='.json':
        data=json.loads(file.read_text(encoding='utf-8-sig'))
        if not isinstance(data,dict) or set(data)!={'version','requirements'} or data['version']!=1: raise ValueError('JSON requires version: 1 and requirements array')
        return data
    if ext in ('.xlsx','.docx'):
        with zipfile.ZipFile(file) as archive:
            if sum(i.file_size for i in archive.infolist())>50*1024*1024: raise ValueError('Expanded document exceeds 50 MiB')
            if ext=='.docx':
                raw=archive.read('word/document.xml')
                if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw: raise ValueError('Unsupported XML declaration')
                root=ET.fromstring(raw);ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                text='\n'.join(''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns))
                result=blocks(text)
        if ext=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(file,read_only=True,data_only=False,keep_links=False)
            try:
                if 'Requirements' not in book.sheetnames: raise ValueError('Workbook needs a Requirements sheet')
                sheet=book['Requirements']
                if sheet.max_row and sheet.max_row>2001: raise ValueError('At most 2000 requirements')
                rows=[]
                for row in sheet.iter_rows():
                    if any(c.data_type=='f' for c in row): raise ValueError('Use literal requirement values, not formulas')
                    rows.append([c.value for c in row])
                result=records(rows)
            finally: book.close()
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(file),on_demand=True)
        try:
            sheet=book.sheet_by_name('Requirements')
            if sheet.nrows>2001: raise ValueError('At most 2000 requirements')
            result=records([sheet.row_values(i) for i in range(sheet.nrows)])
        finally: book.release_resources()
    elif ext=='.doc':
        # Native macOS converter or antiword; never execute Word or document macros.
        if sys.platform=='darwin' and pathlib.Path('/usr/bin/textutil').exists():
            command=['/usr/bin/textutil','-convert','txt','-stdout',str(file)]
        elif shutil.which('antiword'): command=[shutil.which('antiword'),str(file)]
        else: raise ValueError('Legacy DOC needs antiword (Linux/Windows) or macOS textutil; alternatively save as DOCX')
        result=blocks(subprocess.run(command,capture_output=True,check=True,timeout=30).stdout.decode('utf-8'))
    else: raise ValueError('Supported files: JSON, XLS, XLSX, DOCX, DOC')
    return {'version':1,'requirements':result}

if __name__=='__main__':
    try: print(json.dumps(load(sys.argv[1]),ensure_ascii=False))
    except ModuleNotFoundError: print('Install python/requirements-import.txt in the configured Python environment',file=sys.stderr);sys.exit(1)
    except (ValueError,KeyError) as e: print(str(e),file=sys.stderr);sys.exit(1)
    except Exception: print('Requirement import failed; verify format, permissions and converter availability',file=sys.stderr);sys.exit(1)
